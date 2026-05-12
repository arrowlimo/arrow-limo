"""Year-end package builder for accountant-ready exports."""

from __future__ import annotations

import csv
import hashlib
import json
import tempfile
import zipfile
from collections.abc import Iterable
from datetime import date, datetime
from pathlib import Path
from typing import Any

from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas

try:
    from openpyxl import Workbook  # type: ignore
except Exception:  # pragma: no cover - optional dependency fallback
    Workbook = None

from .engine import ensure_audit_storage, generate_audit_check_report
from .schemas import (
    AuditCheckRequest,
    PackageArtifact,
    PackageManifest,
)


def _table_exists(conn, table_name: str) -> bool:
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT 1
            FROM information_schema.tables
            WHERE table_schema = 'public' AND table_name = %s
            LIMIT 1
            """,
            (table_name,),
        )
        return cur.fetchone() is not None


def _has_column(conn, table_name: str, column_name: str) -> bool:
        with conn.cursor() as cur:
                cur.execute(
                        """
                        SELECT 1
                        FROM information_schema.columns
                        WHERE table_schema = 'public'
                            AND table_name = %s
                            AND column_name = %s
                        LIMIT 1
                        """,
                        (table_name, column_name),
                )
                return cur.fetchone() is not None


def _safe_query(conn, sql: str, params: tuple[Any, ...] = ()) -> tuple[list[str], list[tuple[Any, ...]]]:
    with conn.cursor() as cur:
        cur.execute(sql, params)
        rows = cur.fetchall() or []
        headers = [desc[0] for desc in (cur.description or [])]
        return headers, rows


def _write_csv(path: Path, headers: list[str], rows: Iterable[Iterable[Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        for row in rows:
            writer.writerow(["" if value is None else value for value in row])


def _write_json(path: Path, payload: Any) -> None:
    path.write_text(json.dumps(payload, indent=2, default=str), encoding="utf-8")


def _write_text(path: Path, text: str) -> None:
    path.write_text(text, encoding="utf-8")


def _write_pdf(path: Path, title: str, lines: list[str]) -> None:
    pdf = canvas.Canvas(str(path), pagesize=LETTER)
    width, height = LETTER
    y = height - 36
    pdf.setTitle(title)
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(36, y, title)
    y -= 18
    pdf.setFont("Helvetica", 9)
    for line in lines:
        if y < 42:
            pdf.showPage()
            y = height - 36
            pdf.setFont("Helvetica", 9)
        pdf.drawString(36, y, line[:115])
        y -= 12
    pdf.save()


def _write_xlsx(path: Path, sheets: dict[str, tuple[list[str], list[tuple[Any, ...]]]]) -> bool:
    if Workbook is None:
        return False
    wb = Workbook()
    first = True
    for sheet_name, (headers, rows) in sheets.items():
        ws = wb.active if first else wb.create_sheet(title=sheet_name[:31])
        ws.title = sheet_name[:31]
        first = False
        ws.append(headers)
        for row in rows:
            ws.append([None if value is None else value for value in row])
    wb.save(path)
    return True


def _query_legacy_section(conn, table: str, columns: list[str], where_clause: str = "", params: tuple[Any, ...] = ()) -> tuple[list[str], list[tuple[Any, ...]]]:
    sql = f"SELECT {', '.join(columns)} FROM {table}"
    if where_clause:
        sql += f" WHERE {where_clause}"
    sql += " ORDER BY 1"
    return _safe_query(conn, sql, params)


def _build_sections(conn, request: AuditCheckRequest) -> dict[str, tuple[list[str], list[tuple[Any, ...]], list[str]]]:
    date_from = request.date_from or date(request.fiscal_year, 1, 1)
    date_to = request.date_to or date(request.fiscal_year, 12, 31)
    sections: dict[str, tuple[list[str], list[tuple[Any, ...]], list[str]]] = {}

    if _table_exists(conn, "charters"):
        headers, rows = _safe_query(
            conn,
            """
            SELECT charter_id, reserve_number, charter_date, status, client_id,
                   vehicle, routing_type, quoted_hours
            FROM charters
            WHERE COALESCE(charter_date, DATE '1900-01-01') BETWEEN %s AND %s
            ORDER BY charter_date, charter_id
            """,
            (date_from, date_to),
        )
        sections["trip_register"] = (headers, rows, ["charters"])

    if _table_exists(conn, "invoices"):
        headers, rows = _safe_query(
            conn,
            """
            SELECT *
            FROM invoices
            ORDER BY 1
            LIMIT 5000
            """,
        )
        sections["invoice_register"] = (headers, rows, ["invoices"])

    if _table_exists(conn, "driver_payroll"):
        headers, rows = _safe_query(
            conn,
            """
            SELECT employee_id, year, pay_period, gross_pay, cpp, ei, tax_withheld
            FROM driver_payroll
            WHERE year = %s
            ORDER BY employee_id, pay_period
            """,
            (request.fiscal_year,),
        )
        sections["driver_earnings"] = (headers, rows, ["driver_payroll"])

    if _table_exists(conn, "payroll_entries"):
        headers, rows = _safe_query(
            conn,
            """
            SELECT *
            FROM payroll_entries
            WHERE EXTRACT(YEAR FROM COALESCE(pay_date, created_at, NOW())) = %s
            ORDER BY 1
            LIMIT 5000
            """,
            (request.fiscal_year,),
        )
        sections["payroll_register"] = (headers, rows, ["payroll_entries"])

    if _table_exists(conn, "cra_pd7a_returns"):
        headers, rows = _safe_query(
            conn,
            """
            SELECT reporting_year, reporting_month, employee_count,
                   total_gross_payroll, cpp_total, ei_total,
                   income_tax_deducted, total_remittance_due,
                   adjusted_remittance, is_submitted, submission_reference,
                   submitted_by, filing_method, notes
            FROM cra_pd7a_returns
            WHERE reporting_year = %s
            ORDER BY reporting_month
            """,
            (request.fiscal_year,),
        )
        sections["remittance_summary"] = (headers, rows, ["cra_pd7a_returns"])

    if _table_exists(conn, "employee_t4_records"):
        headers, rows = _safe_query(
            conn,
            """
            SELECT *
            FROM employee_t4_records
            WHERE tax_year = %s
            ORDER BY employee_id
            LIMIT 5000
            """,
            (request.fiscal_year,),
        )
        sections["t4_summary"] = (headers, rows, ["employee_t4_records"])
    elif _table_exists(conn, "t4_entries"):
        headers, rows = _safe_query(
            conn,
            """
            SELECT *
            FROM t4_entries
            WHERE tax_year = %s
            ORDER BY employee_id
            LIMIT 5000
            """,
            (request.fiscal_year,),
        )
        sections["t4_summary"] = (headers, rows, ["t4_entries"])

    if _table_exists(conn, "general_ledger"):
        headers, rows = _safe_query(
            conn,
            """
            SELECT *
            FROM general_ledger
            WHERE EXTRACT(YEAR FROM COALESCE(date, created_at, NOW())) = %s
            ORDER BY 1
            LIMIT 5000
            """,
            (request.fiscal_year,),
        )
        sections["general_ledger"] = (headers, rows, ["general_ledger"])

        if _has_column(conn, "general_ledger", "gifi_code") and _has_column(
            conn, "general_ledger", "debit"
        ) and _has_column(conn, "general_ledger", "credit"):
            headers, rows = _safe_query(
                conn,
                """
                SELECT
                    COALESCE(gifi_code::text, 'UNMAPPED') AS gifi_code,
                    COALESCE(account_name, name, 'Unmapped Account') AS account_name,
                    COALESCE(SUM(COALESCE(debit, 0) - COALESCE(credit, 0)), 0) AS net_amount
                FROM general_ledger
                WHERE EXTRACT(YEAR FROM COALESCE(date, created_at, NOW())) = %s
                GROUP BY 1, 2
                ORDER BY 1, 2
                """,
                (request.fiscal_year,),
            )
            sections["gifi_mapped_trial_balance"] = (
                headers,
                rows,
                ["general_ledger"],
            )

    if _table_exists(conn, "audit_events"):
        headers, rows = _safe_query(
            conn,
            """
            SELECT event_id, occurred_at, module, entity_type, entity_id,
                   action, source, correlation_id, retention_until,
                   note, prev_hash, event_hash
            FROM audit_events
            WHERE occurred_at::date BETWEEN %s AND %s
            ORDER BY occurred_at, audit_event_pk
            LIMIT 100000
            """,
            (date_from, date_to),
        )
        sections["audit_events"] = (headers, rows, ["audit_events"])

    if _table_exists(conn, "employee_roe_records"):
        year_col = None
        for candidate in ["tax_year", "year", "reporting_year"]:
            if _has_column(conn, "employee_roe_records", candidate):
                year_col = candidate
                break

        headers, rows = _safe_query(
            conn,
            (
                f"SELECT * FROM employee_roe_records WHERE {year_col} = %s "
                "ORDER BY employee_id LIMIT 5000"
                if year_col
                else "SELECT * FROM employee_roe_records ORDER BY employee_id LIMIT 5000"
            ),
            ((request.fiscal_year,) if year_col else ()),
        )
        sections["roe_summary"] = (headers, rows, ["employee_roe_records"])

    return sections


def generate_notes_to_auditor(
    conn,
    request: AuditCheckRequest,
    report: Any,
    sections: dict[str, tuple[list[str], list[tuple[Any, ...]], list[str]]] | None = None,
) -> str:
    if sections is None:
        sections = _build_sections(conn, request)
    lines: list[str] = []
    lines.append("NOTES TO AUDITOR / ACCOUNTANT")
    lines.append("")
    lines.append(f"System: limousine dispatch, payroll, invoicing, and reporting platform")
    lines.append(f"Fiscal year: {request.fiscal_year}")
    lines.append(f"Generated at: {datetime.utcnow().isoformat()}Z")
    lines.append(f"Generated by: {request.generated_by_name or request.generated_by}")
    lines.append(f"Package mode: {request.package_mode}")
    lines.append("")
    lines.append("System overview")
    lines.append("- Booking/Trip -> Dispatch -> Driver Pay -> Invoicing -> Financial Reporting")
    lines.append("- Employee onboarding -> Payroll calc -> Remittances -> T4/T4 Summary")
    lines.append("")
    lines.append("Data lineage summary")
    for section_name, (_, _, sources) in sections.items():
        lines.append(f"- {section_name}: source tables = {', '.join(sources)}")
    lines.append("")
    lines.append("Reconciliations performed")
    for finding in getattr(report, "findings", []):
        lines.append(f"- {finding.check_id}: {finding.status} - {finding.message}")
    lines.append("")
    lines.append("Exceptions / anomalies")
    exceptions = [f for f in getattr(report, "findings", []) if f.status != "PASS"]
    if exceptions:
        for finding in exceptions:
            lines.append(f"- {finding.check_id}: {finding.message}")
    else:
        lines.append("- None detected by automated checks.")
    lines.append("")
    lines.append("Questions / unknowns")
    unknowns = [f for f in getattr(report, "findings", []) if getattr(f, "requires_confirmation", False)]
    if unknowns:
        for finding in unknowns:
            lines.append(f"- {finding.check_id}: requires confirmation")
    else:
        lines.append("- None")
    lines.append("")
    lines.append("Retention")
    lines.append("- Generated package and source records are intended to be retained for 6+ years (general rule; confirm with your retention policy).")
    lines.append("- Audit trail is append-only and hash chained.")
    return "\n".join(lines)


def generate_year_end_package(
    conn, request: AuditCheckRequest, output_dir: str | Path | None = None
) -> PackageManifest:
    ensure_audit_storage(conn)
    report = generate_audit_check_report(conn, request)
    sections = _build_sections(conn, request)
    package_id = datetime.utcnow().strftime("%Y%m%d%H%M%S") + "_" + request.fiscal_year.__str__()
    package_name = request.package_name or f"year_end_package_{request.fiscal_year}"
    root = Path(output_dir) if output_dir else Path(tempfile.mkdtemp(prefix="audit_package_"))
    package_dir = root / package_id
    package_dir.mkdir(parents=True, exist_ok=True)

    artifacts: list[PackageArtifact] = []

    manifest_payload = {
        "package_id": package_id,
        "package_name": package_name,
        "fiscal_year": request.fiscal_year,
        "date_from": request.date_from.isoformat() if request.date_from else None,
        "date_to": request.date_to.isoformat() if request.date_to else None,
        "generated_by": request.generated_by,
        "generated_by_name": request.generated_by_name,
        "correlation_id": request.correlation_id,
        "package_mode": request.package_mode,
        "retention_policy": "6+ years",
        "overall_status": report.overall_status,
        "sections": list(sections.keys()),
    }

    for section_name, (headers, rows, sources) in sections.items():
        csv_path = package_dir / f"{section_name}.csv"
        json_path = package_dir / f"{section_name}.json"
        xlsx_path = package_dir / f"{section_name}.xlsx"
        pdf_path = package_dir / f"{section_name}.pdf"
        _write_csv(csv_path, headers, rows)
        _write_json(json_path, {"headers": headers, "rows": rows})
        _write_pdf(pdf_path, f"{section_name} ({request.fiscal_year})", [
            f"Rows: {len(rows)}",
            f"Sources: {', '.join(sources)}",
            "See CSV/JSON/XLSX artifacts for full detail.",
        ])
        xlsx_ok = _write_xlsx(xlsx_path, {section_name: (headers, rows)})
        if not xlsx_ok:
            xlsx_path.write_text("XLSX generation requires openpyxl to be installed.", encoding="utf-8")

        for artifact_path, fmt in [
            (csv_path, "csv"),
            (json_path, "json"),
            (xlsx_path, "xlsx"),
            (pdf_path, "pdf"),
        ]:
            artifacts.append(
                PackageArtifact(
                    section=section_name,
                    file_name=artifact_path.name,
                    format=fmt,  # type: ignore[arg-type]
                    sha256=hashlib.sha256(artifact_path.read_bytes()).hexdigest(),
                    size_bytes=artifact_path.stat().st_size,
                    source_tables=sources,
                )
            )

    checks_json = {
        "fiscal_year": request.fiscal_year,
        "generated_at": report.generated_at.isoformat(),
        "overall_status": report.overall_status,
        "summary": report.summary,
        "findings": [finding.model_dump(mode="json") for finding in report.findings],
    }
    checks_path = package_dir / "audit_checks.json"
    _write_json(checks_path, checks_json)

    notes_text = generate_notes_to_auditor(conn, request, report, sections)
    notes_path = package_dir / "notes_to_auditor.md"
    _write_text(notes_path, notes_text)
    notes_pdf_path = package_dir / "notes_to_auditor.pdf"
    _write_pdf(notes_pdf_path, "Notes to Auditor", notes_text.splitlines())
    notes_email_path = package_dir / "email_draft.txt"
    _write_text(
        notes_email_path,
        "Subject: Year-End / CRA Package Ready\n\n" + notes_text,
    )

    manifest_path = package_dir / "manifest.json"
    zip_path = package_dir / f"{package_name}_{package_id}.zip"
    _write_json(
        manifest_path,
        {
            **manifest_payload,
            "artifacts": [artifact.model_dump(mode="json") for artifact in artifacts],
            "checks": [finding.model_dump(mode="json") for finding in report.findings],
            "notes_file": notes_path.name,
        },
    )

    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for path in package_dir.iterdir():
            if path == zip_path:
                continue
            zf.write(path, arcname=path.name)

    content_hash = hashlib.sha256(zip_path.read_bytes()).hexdigest()
    manifest = PackageManifest(
        package_id=package_id,
        package_name=package_name,
        package_mode=request.package_mode,
        fiscal_year=request.fiscal_year,
        date_from=request.date_from,
        date_to=request.date_to,
        generated_by=request.generated_by,
        generated_by_name=request.generated_by_name,
        correlation_id=request.correlation_id,
        overall_status=report.overall_status,
        content_hash=content_hash,
        zip_path=str(zip_path),
        notes_path=str(notes_path),
        artifacts=artifacts,
        checks=report.findings,
    )

    with conn.cursor() as cur:
        cur.execute(
            """
            INSERT INTO audit_package_runs (
                package_id, package_name, package_mode, fiscal_year,
                date_from, date_to, generated_by, generated_by_name,
                correlation_id, overall_status, content_hash, zip_path,
                notes_path, manifest_json, checks_json, status
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s::jsonb, %s::jsonb, %s)
            ON CONFLICT (package_id) DO UPDATE
            SET content_hash = EXCLUDED.content_hash,
                zip_path = EXCLUDED.zip_path,
                notes_path = EXCLUDED.notes_path,
                manifest_json = EXCLUDED.manifest_json,
                checks_json = EXCLUDED.checks_json,
                status = EXCLUDED.status
            """,
            (
                manifest.package_id,
                manifest.package_name,
                manifest.package_mode,
                manifest.fiscal_year,
                manifest.date_from,
                manifest.date_to,
                manifest.generated_by,
                manifest.generated_by_name,
                manifest.correlation_id,
                manifest.overall_status,
                manifest.content_hash,
                manifest.zip_path,
                manifest.notes_path,
                json.dumps(manifest_payload),
                json.dumps(checks_json),
                "ready",
            ),
        )
    conn.commit()
    return manifest
